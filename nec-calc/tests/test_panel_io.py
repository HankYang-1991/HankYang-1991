import math
import os
import re
import shutil
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from necalc import engine as E                 # noqa: E402
from necalc import loads as L                  # noqa: E402
from necalc import panel_io as P               # noqa: E402
from necalc import report as R                 # noqa: E402
from necalc import xlsx as X                   # noqa: E402
from necalc.i18n import TEXT, tr               # noqa: E402


def _wiring(p):
    return E.WiringInput(voltage=208, wiring="3ph4w",
                         length_ft=p.feeder_length_ft or 100)


def _loads(p):
    return [vars(x) for x in p.loads]


class PanelIOTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def path(self, name):
        return os.path.join(self.dir, name)

    # ---- round trips -------------------------------------------------
    def test_template_round_trip(self):
        for lang in ("en", "zh"):
            f = self.path(f"tpl_{lang}.xlsx")
            P.write_template(f, lang)
            sheets = X.read_xlsx(f)
            self.assertEqual(list(sheets), [tr("xl_sh_instr", lang),
                                            tr("xl_sh_panels", lang),
                                            tr("xl_sh_loads", lang)])
            panels, notes = P.read_panels(f)
            self.assertEqual(notes, [])
            ex = P.example_panels()
            self.assertEqual([p.name for p in panels], ["LP-1", "DP-1"])
            for a, b in zip(panels, ex):
                self.assertEqual(_loads(a), _loads(b))
                self.assertEqual((a.system, a.diversity, a.target_pf,
                                  a.feeder_length_ft),
                                 (b.system, b.diversity, b.target_pf,
                                  b.feeder_length_ft))

    def test_results_workbook(self):
        panels = P.example_panels()
        res = P.calculate_all(panels, "2023", _wiring)
        self.assertTrue(all(it["calc"] and it["feeder"] for it in res))
        for lang in ("en", "zh"):
            f = self.path(f"res_{lang}.xlsx")
            P.write_results(f, panels, res, lang)
            sheets = X.read_xlsx(f)
            self.assertEqual(list(sheets)[:3], [tr("xl_sh_summary", lang),
                                                tr("xl_sh_panels", lang),
                                                tr("xl_sh_loads", lang)])
            self.assertIn("LP-1", sheets)
            # re-import: the Loads sheet wins over the schedules
            p2, notes = P.read_panels(f)
            self.assertEqual(notes, [])
            for a, b in zip(panels, p2):
                self.assertEqual(_loads(a), _loads(b))
            # totals: SUM formulas whose cached value is the exact sum
            summ = sheets[tr("xl_sh_summary", lang)]
            hdr = next(i for i, r in enumerate(summ)
                       if r[0] == tr("xl_panel", lang))
            col = summ[hdr].index(tr("xl_conn_kva", lang))
            total = next(r for r in summ if r[0] == tr("xl_total", lang))
            self.assertAlmostEqual(total[col], math.fsum(
                it["calc"]["connected"]["kva_arith"] for it in res), 9)
            with zipfile.ZipFile(f) as z:
                xml = z.read("xl/worksheets/sheet1.xml").decode()
            self.assertIn("<f>SUM(F5:F6)</f>", xml)

    def test_summary_text(self):
        res = P.calculate_all(P.example_panels(), "2023", _wiring)
        for lang in ("en", "zh"):
            lines, errors = P.summary_text_lines(res, lang)
            self.assertEqual(errors, [])
            self.assertTrue(any(ln.startswith("LP-1") for ln in lines))
            self.assertTrue(lines[-1].startswith(tr("xl_total", lang)))
            self.assertIn("<table>", P.summary_html(res, lang))
            self.assertIn("DP-1", P.summary_csv(res, lang))

    def test_errors_in_results(self):
        bad = L.Panel("BAD", "208Y/120",
                      loads=[L.Load("x", unit="HP", rating=3, voltage=999,
                                    phases=3, hp="3")])
        bad.loads[0].unit = "??"
        res = P.calculate_all([bad] + P.example_panels(), "2023", _wiring)
        self.assertIsNotNone(res[0]["error"])
        f = self.path("err.xlsx")
        P.write_results(f, [it["panel"] for it in res], res, "en")
        summ = X.read_xlsx(f)["Summary"]
        row = next(r for r in summ if r[0] == "BAD")
        self.assertIn("??", row[-1])

    # ---- tolerant reader --------------------------------------------
    def test_messy_chinese_list(self):
        rows = [
            ["配電盤負載表"],
            [],
            ["盤名", "系統", "迴路", "負載名稱", "類別", "台數", "容量", "單位",
             "功因", "電壓", "相", "連續", "需量因數", "幹線長度"],
            ["LP-2F", "120/208V", 1, "辦公室照明", "照明", 10, 400, "va",
             "95%", 120, "R", "是", None, "30m"],
            [None, None, 3, "插座", "插座", 40, 180, None, 0.9, None, "S",
             "否", "自動", None],
            [None, None, 5, "排風機", "馬達", 2, "1-1/2", "HP", None, 208, 3,
             None, None, None],
            [None, None, 7, "冷氣", "冷氣", 1, 5, "kW", 85, 208, "2", "N",
             "80%", None],
            [None, None, None, "小計", None, None, 9999],
            ["MDP", "480/277", 1, "Chiller", "hvac", 1, 150, "kVA", 0.9, 480,
             "3Ø", "Y", None, None],
            [None, None, 2, "Bad row", "lighting", 1, "abc", "VA"],
            [None, None, 3, "Weird", "xyz", 1, 100, "VA", 2.5, None, "Q",
             "maybe"],
        ]
        panels, notes = P.parse_sheets({"負載表": rows})
        lp, mdp = panels
        self.assertEqual((lp.name, lp.system), ("LP-2F", "208Y/120"))
        self.assertAlmostEqual(lp.feeder_length_ft, 30 * P.M_TO_FT)
        self.assertEqual(len(lp.loads), 4)            # subtotal skipped
        a, b, c, d = lp.loads
        self.assertEqual((a.category, a.pf, a.continuous, a.phase_conn),
                         ("lighting", 0.95, True, "A"))
        self.assertEqual((b.voltage, b.phase_conn, b.demand_factor),
                         (120.0, "B", None))
        self.assertEqual((c.rating, c.unit, c.pf, c.phases),
                         (1.5, "HP", P.DEFAULT_MOTOR_PF, 3))
        self.assertEqual((d.category, d.pf, d.phases, d.demand_factor),
                         ("cooling", 0.85, 1, 0.8))
        self.assertEqual((mdp.system, len(mdp.loads)), ("480Y/277", 2))
        weird = mdp.loads[1]
        self.assertEqual((weird.category, weird.pf, weird.voltage),
                         ("other", 1.0, 277.0))
        keys = [k for k, _ in notes]
        for k in ("io_m_to_ft", "io_skip_rating", "io_bad_cat",
                  "io_bad_num", "io_bad_conn", "io_bad_bool"):
            self.assertIn(k, keys)
        for lang in ("en", "zh"):
            for line in P.issue_lines(notes, lang):
                self.assertNotIn("{", line)
                self.assertNotIn("xl_", line)

    def test_panels_sheet_and_defaults(self):
        sheets = {
            "Panels": [["Panel", "System", "Diversity", "Target PF",
                        "Feeder length (m)"],
                       ["P1", "480", 0.8, "90%", 50],
                       ["P1", "208Y/120"],
                       ["P2", "600V"]],
            "Loads": [["Description", "Load (kW)", "Ph"],
                      ["Heater", 10, 3]],
        }
        panels, notes = P.parse_sheets(sheets)
        p1, p2 = panels
        self.assertEqual(p1.system, "480D")
        self.assertAlmostEqual(p1.diversity, 1.25)
        self.assertAlmostEqual(p1.target_pf, 0.9)
        self.assertAlmostEqual(p1.feeder_length_ft, 50 * P.M_TO_FT)
        self.assertEqual(p1.loads[0].unit, "kW")
        self.assertEqual(p1.loads[0].voltage, 480.0)
        self.assertEqual(p2.system, "208Y/120")       # inferred
        keys = [k for k, _ in notes]
        for k in ("io_coincidence", "io_dup_panel", "io_bad_system",
                  "io_new_panel", "io_no_panel_col"):
            self.assertIn(k, keys)

    def test_old_csv_export(self):
        p = L.Panel("LP-1", "208Y/120",
                    loads=[L.Load(**vars(x)) for x in L.EXAMPLE_LOADS])
        c = L.calculate(p)
        for lang in ("en", "zh"):
            f = self.path(f"LP-1_{lang}.csv")
            with open(f, "w", encoding="utf-8-sig", newline="") as fh:
                fh.write(R.load_calc_csv(p, c, lang))
            panels, notes = P.read_panels(f)
            self.assertEqual([(q.name, q.system) for q in panels],
                             [("LP-1", "208Y/120")])
            self.assertEqual([k for k, _ in notes], ["io_computed"])
            for a, b in zip(p.loads, panels[0].loads):
                da, db = vars(a), dict(vars(b), circuit=a.circuit)
                self.assertEqual(da, db)

    def test_csv_big5_semicolon(self):
        f = self.path("list.csv")
        text = ("盤名;負載名稱;類別;數量;額定;單位;電壓;相數\r\n"
                "LP-A;照明;照明;5;1,200;VA;120;1\r\n"
                "LP-A;馬達;馬達;1;5HP;;208;3\r\n")
        with open(f, "wb") as fh:
            fh.write(text.encode("cp950"))
        panels, notes = P.read_panels(f)
        self.assertEqual(len(panels), 1)
        lt, mt = panels[0].loads
        self.assertEqual((lt.qty, lt.rating, lt.unit), (5, 1200.0, "VA"))
        self.assertEqual((mt.rating, mt.unit, mt.phases), (5.0, "HP", 3))
        self.assertEqual([k for k, _ in notes], ["io_new_panel"])

    def test_shared_strings_reader(self):
        """Workbooks saved by Excel use a shared string table."""
        f = self.path("excel.xlsx")
        ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        rns = ('xmlns:r="http://schemas.openxmlformats.org/officeDocument/'
               '2006/relationships"')
        strings = ["Panel", "Description", "Rating", "LP-9", "Fan",
                   ("<r><t>Ki</t></r><r><rPr><b/></rPr><t>tchen</t></r>"
                    "<rPh sb=\"0\" eb=\"1\"><t>X</t></rPh>")]
        sst = "".join(f"<si>{s}</si>" if s.startswith("<r>")
                      else f"<si><t>{s}</t></si>" for s in strings)
        sheet = (f'<worksheet {ns}><sheetData>'
                 '<row r="1"><c r="A1" t="s"><v>0</v></c>'
                 '<c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c>'
                 '</row><row r="2"><c r="A2" t="s"><v>3</v></c>'
                 '<c r="B2" t="s"><v>4</v></c><c r="C2"><v>1500</v></c></row>'
                 '<row r="4"><c r="B4" t="s"><v>5</v></c>'
                 '<c r="C4"><f>1000*2</f><v>2000</v></c></row>'
                 '</sheetData></worksheet>')
        with zipfile.ZipFile(f, "w") as z:
            z.writestr("[Content_Types].xml", "<Types/>")
            z.writestr("_rels/.rels",
                       '<Relationships xmlns="http://schemas.openxmlformats.'
                       'org/package/2006/relationships"><Relationship '
                       'Id="rId1" Type="http://schemas.openxmlformats.org/'
                       'officeDocument/2006/relationships/officeDocument" '
                       'Target="xl/workbook.xml"/></Relationships>')
            z.writestr("xl/workbook.xml",
                       f'<workbook {ns} {rns}><sheets><sheet name="Sheet1" '
                       'sheetId="1" r:id="rId1"/></sheets></workbook>')
            z.writestr("xl/_rels/workbook.xml.rels",
                       '<Relationships xmlns="http://schemas.openxmlformats.'
                       'org/package/2006/relationships">'
                       '<Relationship Id="rId1" Type="http://schemas.'
                       'openxmlformats.org/officeDocument/2006/relationships/'
                       'worksheet" Target="worksheets/sheet1.xml"/>'
                       '<Relationship Id="rId2" Type="http://schemas.'
                       'openxmlformats.org/officeDocument/2006/relationships/'
                       'sharedStrings" Target="sharedStrings.xml"/>'
                       '</Relationships>')
            z.writestr("xl/sharedStrings.xml", f"<sst {ns}>{sst}</sst>")
            z.writestr("xl/worksheets/sheet1.xml", sheet)
        rows = X.read_xlsx(f)["Sheet1"]
        self.assertEqual(rows[1], ["LP-9", "Fan", 1500.0])
        self.assertEqual(rows[3], [None, "Kitchen", 2000.0])
        panels, _ = P.read_panels(f)
        self.assertEqual([ld.name for ld in panels[0].loads],
                         ["Fan", "Kitchen"])
        self.assertEqual(panels[0].name, "LP-9")

    def test_read_errors(self):
        with self.assertRaises(E.CalcError) as cm:
            P.read_panels(self.path("old.xls"))
        self.assertEqual(cm.exception.key, "err_xls")
        f = self.path("broken.xlsx")
        with open(f, "wb") as fh:
            fh.write(b"not a zip")
        with self.assertRaises(E.CalcError) as cm:
            P.read_panels(f)
        self.assertEqual(cm.exception.key, "err_import")
        f = self.path("empty.csv")
        with open(f, "w") as fh:
            fh.write("a,b,c\n1,2,3\n")
        with self.assertRaises(E.CalcError) as cm:
            P.read_panels(f)
        self.assertEqual(cm.exception.key, "io_no_data")

    def test_value_parsers(self):
        self.assertEqual(P._num("1-1/2")[0], 1.5)
        self.assertEqual(P._num("7.5 hp"), (7.5, False, "hp"))
        self.assertEqual(P._num("85%"), (85.0, True, ""))
        self.assertEqual(P._num("1,200")[0], 1200.0)
        with self.assertRaises(ValueError):
            P._num("abc")
        self.assertEqual(P._frac(85, False), 0.85)
        self.assertEqual(P._frac(0.9, False), 0.9)
        with self.assertRaises(ValueError):
            P._frac(2.5, False)
        for raw, code in (("120/208Y", "208Y/120"), ("480Y/277V", "480Y/277"),
                          ("480V 3W", "480D"), ("480Δ", "480D"),
                          ("120/240V 1PH", "240/120 1PH"),
                          ("120V", "120 1PH"), ("600V", None)):
            self.assertEqual(P.system_code(raw), code, raw)
        for raw, code in (("A-B", "AB"), ("ac", "CA"), ("T", "C"),
                          ("L1 L2 L3", "ABC"), ("自動", "auto"), ("", "auto"),
                          ("X", None)):
            self.assertEqual(P._phase_conn(raw), code, raw)
        self.assertEqual(P._phases("3Ø"), (3, False))
        self.assertEqual(P._phases(2.0), (1, True))
        self.assertIsNone(P._phases("4"))

    def test_header_matching(self):
        for h, f in (("Rating (per unit)", "rating"), ("Load (VA)", "rating:VA"),
                     ("Connected kW", "_ignore"), ("Control", None),
                     ("Demand factor (blank = auto)", "demand_factor"),
                     ("需量因數 (空白=自動)", "demand_factor"),
                     ("Phase connection", "phase_conn"),
                     ("Phase A (A)", "_ignore"), ("Phases", "phases"),
                     ("Voltage drop %", "_ignore"), ("V", "voltage"),
                     ("幹線長度 (ft，空白=依程式設定)", "feeder_length")):
            self.assertEqual(P.match_header(h), f, h)

    def test_text_keys(self):
        keys = [k for k in TEXT if k.startswith(("xl_", "io_", "ps_"))]
        self.assertGreater(len(keys), 80)
        for k in keys:
            en, zh = TEXT[k]
            self.assertEqual(sorted(re.findall(r"{(\w+)}", en)),
                             sorted(re.findall(r"{(\w+)}", zh)), k)


class OpenpyxlCompatTest(unittest.TestCase):
    """Optional: files written by openpyxl (shared strings) read back, and
    our files load in openpyxl."""

    @classmethod
    def setUpClass(cls):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            raise unittest.SkipTest("openpyxl not installed")

    def test_openpyxl(self):
        import openpyxl
        d = tempfile.mkdtemp()
        try:
            f = os.path.join(d, "o.xlsx")
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "負載清單"
            ws.append(["盤名", "負載名稱", "額定", "單位"])
            ws.append(["LP-1", "照明", 600, "VA"])
            wb.save(f)
            panels, _ = P.read_panels(f)
            self.assertEqual(panels[0].loads[0].rating, 600.0)
            g = os.path.join(d, "t.xlsx")
            P.write_template(g, "zh")
            wb = openpyxl.load_workbook(g)
            self.assertEqual(wb.sheetnames, ["說明", "盤清單", "負載清單"])
        finally:
            shutil.rmtree(d, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
